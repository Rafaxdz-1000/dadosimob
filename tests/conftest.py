from datetime import datetime

import pytest
from openpyxl import Workbook

HEADERS = [
    "N° do Cadastro (SQL)", "Nome do Logradouro", "Número", "Complemento", "Bairro", "Referência",
    "CEP", "Natureza de Transação", "Valor de Transação (declarado pelo contribuinte)",
    "Data de Transação", "Valor Venal de Referência", "Proporção Transmitida (%)",
    "Valor Venal de Referência (proporcional)", "Base de Cálculo adotada", "Tipo de Financiamento",
    "Valor Financiado", "Cartório de Registro", "Matrícula do Imóvel", "Situação do SQL",
    "Área do Terreno (m2)", "Testada (m)", "Fração Ideal", "Área Construída (m2)", "Uso (IPTU)",
    "Descrição do uso (IPTU)", "Padrão (IPTU)", "Descrição do padrão (IPTU)", "ACC (IPTU)",
]


def _row(sql, rua, valor, data, area_c, area_t, uso, natureza="1.Compra e venda", proporcao=100):
    venal = 400000.0
    return [
        sql, rua, 100, None, "MOOCA", None, 3104000, natureza, valor, data, venal, proporcao,
        venal, valor, "Nenhum", 0, 7, 12345, "Normal", area_t, 10, 0.01, area_c, 64, uso,
        "2B", "RESIDENCIAL VERTICAL - PADRÃO B", 2010,
    ]


def _save(path, sheets):
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets:
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    wb.save(path)
    return path


@pytest.fixture
def workbook(tmp_path):
    """Workbook shaped like the official file: explanation sheet + monthly sheets."""
    wb = Workbook()
    info = wb.active
    info.title = "EXPLICAÇÕES"
    info.append(["Este arquivo traz os dados das DTIs pagas."])

    jan = wb.create_sheet("JAN-2024")
    jan.append(HEADERS)
    jan.append(_row(1000100011, "R JUVENTUS", 500000.0, datetime(2023, 12, 20), 50, 1000,
                    "APARTAMENTO EM CONDOMÍNIO (EXCETO VAGA)"))
    padded_row = _row(1000100022, " R TAQUARI ", "1.200.000,00", "15/01/2024", 150, 200,
                      " RESIDÊNCIA ")
    padded_row[3] = "   "  # A text field containing only spaces should be missing.
    padded_row[4] = " MOOCA "
    padded_row[7] = " 1.Compra e venda "
    padded_row[26] = " RESIDENCIAL VERTICAL - PADRÃO B "
    jan.append(padded_row)
    jan.append(_row(1000100033, "R ORATORIO", 300000.0, datetime(2024, 1, 5), 0, 250, "TERRENO"))
    jan.append(_row(1000100044, "R ERRO", 1_000.0, datetime(2024, 1, 5), 40, 100, "LOJA"))  # too cheap
    jan.append([None] * len(HEADERS))

    fev = wb.create_sheet("FEV-2024")
    fev.append(["Tabela de fevereiro"])  # title row above the header
    fev.append(HEADERS)
    fev.append(_row(1000100055, "R DA MOOCA", 800000.0, datetime(2024, 2, 1), 80, 1000,
                    "APARTAMENTO EM CONDOMÍNIO (EXCETO VAGA)", natureza="3.Doação"))
    fev.append(_row(1000100066, "R X", 900000.0, datetime(2024, 2, 1), 10, 100,
                    "VAGA DE GARAGEM EM CONDOMÍNIO"))  # 90k/m²: plausible cap test

    usos = wb.create_sheet("Tabela de USOS")
    usos.append(["Código", "Descrição"])
    usos.append([64, "APARTAMENTO"])

    path = tmp_path / "itbi_2024.xlsx"
    wb.save(path)
    return path


@pytest.fixture
def workbook_2019(tmp_path):
    """2019 to 2022: the "Descrição do padrão" column is labeled as a second "ACC (IPTU)"."""
    headers = ["ACC (IPTU)" if h == "Descrição do padrão (IPTU)" else h for h in HEADERS]
    return _save(tmp_path / "itbi_2019.xlsx", [
        ("JAN-2019", [headers, _row(1000100011, "R JUVENTUS", 500000.0, datetime(2019, 1, 10), 50, 1000,
                                    "APARTAMENTO EM CONDOMÍNIO (EXCETO VAGA)")]),
    ])


@pytest.fixture
def workbook_typo(tmp_path):
    """FEV-2026 of the official file spells the header as "Descrição do pardão (IPTU)"."""
    headers = ["Descrição do pardão (IPTU)" if h == "Descrição do padrão (IPTU)" else h for h in HEADERS]
    return _save(tmp_path / "itbi_2026.xlsx", [
        ("FEV-2026", [headers, _row(1000100011, "R JUVENTUS", 500000.0, datetime(2026, 2, 10), 50, 1000,
                                    "APARTAMENTO EM CONDOMÍNIO (EXCETO VAGA)")]),
    ])


@pytest.fixture
def edge_workbook(tmp_path):
    """Quirks found in the official 2024 file."""
    return _save(tmp_path / "itbi_edge.xlsx", [
        # JAN-2024 and OUT-2024 of the official file have no header row: data starts at row 0.
        ("JAN-2024", [
            _row(1000100011, "R JUVENTUS", 500000.0, datetime(2024, 1, 10), 50, 1000,
                 "APARTAMENTO EM CONDOMÍNIO (EXCETO VAGA)"),
            # A unit sold on the parent lot of a new building: a 0.25% share of the land.
            _row(1000100099, "R DO LOTE MAE", 400000.0, datetime(2024, 1, 12), 0, 8000,
                 "TERRENO", proporcao=0.25),
        ]),
        ("FEV-2024", [HEADERS, _row(1000100022, "R TAQUARI", 600000.0, datetime(2024, 2, 1), 60, 200,
                                    "RESIDÊNCIA")]),
        ("MAR-2024", [["resumo", 1, 2], ["total", 3, 4]]),  # monthly name, unknown layout
        ("ABR-2024", [["x"] * len(HEADERS), ["y"] * len(HEADERS)]),  # right width, no numbers
        ("Tabela de USOS", [["Código", "Descrição"], [64, "APARTAMENTO"]]),
    ])
