import logging
from datetime import date

import pandas as pd
import pytest
import requests

from dadosimob._text import parse_br_date, parse_br_number
from dadosimob.itbi import sp


def test_read_workbook(workbook):
    df = sp.read(workbook)
    assert len(df) == 6  # blank row and non-transaction sheets are skipped
    assert set(df["mes_referencia"].dt.month) == {1, 2}
    row = df.iloc[1]
    assert row["valor_transacao"] == 1_200_000.0
    assert row["data_transacao"] == pd.Timestamp("2024-01-15")
    assert (df["codigo_ibge"] == "3550308").all()


def test_codes_keep_leading_zeros(workbook):
    row = sp.read(workbook).iloc[1]
    assert row["sql"] == "01000100022"  # the SQL has 11 digits
    assert row["cep"] == "03104000"  # every CEP in the city starts with 0


def test_sheet_without_header_reuses_the_header_of_other_sheets(edge_workbook):
    df = sp.read(edge_workbook)
    jan = df[df["mes_referencia"] == pd.Timestamp("2024-01-01")]
    assert len(jan) == 2
    assert jan.iloc[0]["logradouro"] == "R JUVENTUS"
    assert jan.iloc[0]["valor_transacao"] == 500_000.0
    assert df.iloc[0]["logradouro"] == "R JUVENTUS"  # sheet order is kept


@pytest.mark.parametrize("sheet", ["MAR-2024", "ABR-2024"])
def test_monthly_sheet_that_cannot_be_read_is_skipped_with_a_warning(edge_workbook, caplog, sheet):
    with caplog.at_level(logging.WARNING, logger="dadosimob"):
        df = sp.read(edge_workbook)
    assert set(df["mes_referencia"].dt.month) == {1, 2}
    assert any(sheet in r.getMessage() for r in caplog.records if r.levelno == logging.WARNING)


def test_unrecognized_header_warns_with_sheet_name(workbook, caplog):
    from openpyxl import load_workbook

    with caplog.at_level(logging.WARNING, logger="dadosimob"):
        sp.read(workbook)
    assert not any("cabeçalhos não reconhecidos" in record.getMessage() for record in caplog.records)

    caplog.clear()
    wb = load_workbook(workbook)
    sheet = wb["JAN-2024"]
    blank_column = sheet.max_column + 1
    sheet.cell(row=2, column=blank_column, value="sem cabeçalho")
    sheet.cell(row=1, column=blank_column + 1, value="Nova coluna da Prefeitura")
    wb.save(workbook)

    with caplog.at_level(logging.WARNING, logger="dadosimob"):
        sp.read(workbook)
    unmapped_warnings = [
        record.getMessage()
        for record in caplog.records
        if record.levelno == logging.WARNING and "cabeçalhos não reconhecidos" in record.getMessage()
    ]
    assert len(unmapped_warnings) == 1
    assert "JAN-2024" in unmapped_warnings[0]
    assert "Nova coluna da Prefeitura" in unmapped_warnings[0]
    assert "nan" not in unmapped_warnings[0].lower()


def test_padrao_description_labeled_as_acc_goes_to_its_own_column(workbook_2019):
    df = sp.read(workbook_2019)
    assert "descricao_padrao" in df.columns
    assert df.loc[0, "descricao_padrao"] == "RESIDENCIAL VERTICAL - PADRÃO B"
    assert df.loc[0, "acc_iptu"] == "2010"  # the real ACC is the last column


def test_misspelled_padrao_header_is_recognized(workbook_typo):
    df = sp.read(workbook_typo)
    assert "descricao_padrao" in df.columns
    assert df.loc[0, "descricao_padrao"] == "RESIDENCIAL VERTICAL - PADRÃO B"


def test_partial_transfer_has_no_price_per_m2(edge_workbook):
    df = sp.read(edge_workbook)
    unit = df[df["logradouro"] == "R DO LOTE MAE"].iloc[0]
    assert unit["proporcao_transmitida"] == 0.25
    assert pd.isna(unit["preco_m2"])  # the file has no area for the unit itself


def test_columns_are_canonical(workbook):
    df = sp.read(workbook)
    for col in sp.COLUMNS:
        assert col in df.columns, col
    assert {"mes_referencia", "tipo_imovel", "preco_m2", "codigo_ibge"} <= set(df.columns)


def test_property_type_and_price_m2(workbook):
    df = sp.read(workbook)
    assert df["tipo_imovel"].tolist() == ["apartamento", "casa", "terreno", "comercial", "apartamento", "garagem"]
    assert df.loc[0, "preco_m2"] == 10_000.0
    assert df.loc[2, "preco_m2"] == 1_200.0  # land uses terrain area


def test_filter_months(workbook):
    df = sp.read(workbook, months=[2])
    assert len(df) == 2
    assert (df["mes_referencia"] == pd.Timestamp("2024-02-01")).all()


def test_clean(workbook):
    df = sp.read(workbook)
    cleaned = sp.clean(df)
    assert "R ERRO" not in cleaned["logradouro"].tolist()
    sales = sp.clean(df, only_sales=True)
    assert "R DA MOOCA" not in sales["logradouro"].tolist()


@pytest.mark.parametrize(
    ("raw", "expected"),
    [("R$ 1.234.567,89", 1234567.89), ("1234,5", 1234.5), ("12,5%", 12.5), ("1.234.567", 1234567.0),
     (1500, 1500.0), ("", None), ("-", None), ("abc", None), (float("nan"), None)],
)
def test_parse_br_number(raw, expected):
    assert parse_br_number(raw) == expected


def test_parse_br_date():
    assert parse_br_date("15/01/2024") == date(2024, 1, 15)
    assert parse_br_date("2024-01-15 00:00:00") == date(2024, 1, 15)
    assert parse_br_date(45306) == date(2024, 1, 15)
    assert parse_br_date("nada") is None


@pytest.mark.parametrize(
    ("name", "expected"),
    [("JAN-2024", date(2024, 1, 1)), ("Março 2023", date(2023, 3, 1)), ("DEZ-2019", date(2019, 12, 1)),
     ("Tabela de USOS", None), ("EXPLICAÇÕES", None)],
)
def test_sheet_period(name, expected):
    assert sp.sheet_period(name) == expected


PAGE = """
<ul>
<li>2026 (<a href="https://prefeitura.sp.gov.br/documents/d/fazenda/guias-de-itbi-pagas-27082026-xls-xlsx">Excel/xlsx</a>)
 (<a href="http://prefeitura.sp.gov.br/x/GUIAS%20DE%20ITBI%20PAGAS%20ODS.ods">ODS</a>)</li>
<li>2024 (<a href="/cidade/secretarias/upload/fazenda/arquivos/itbi/GUIAS-DE-ITBI-PAGAS-2024.xlsx">Excel/xlsx</a>)</li>
<li>2019 (<a href="/cidade/secretarias/upload/fazenda/arquivos/itbi/ITBI_Setembro_2022/GUIAS_DE_ITBI_PAGAS_(2019).xlsx">Excel/xlsx</a>)</li>
</ul>
"""


def test_parse_source_page():
    files = sp.parse_source_page(PAGE)
    xlsx = {f.year: f.url for f in files if f.format == "xlsx"}
    assert set(xlsx) == {2019, 2024, 2026}
    assert xlsx[2024] == "https://prefeitura.sp.gov.br/cidade/secretarias/upload/fazenda/arquivos/itbi/GUIAS-DE-ITBI-PAGAS-2024.xlsx"
    assert xlsx[2026].endswith("27082026-xls-xlsx")
    assert xlsx[2019].endswith("GUIAS_DE_ITBI_PAGAS_(2019).xlsx")  # year in folder name must not confuse


def test_download_retries_when_the_page_misses_the_year(monkeypatch, tmp_path):
    calls = []

    def flaky_list_files():
        calls.append(1)
        return [] if len(calls) == 1 else [sp.ItbiFile(2024, "https://example.org/2024.xlsx", "xlsx")]

    monkeypatch.setattr(sp, "list_files", flaky_list_files)
    monkeypatch.setattr(sp._http, "download", lambda url, **kw: tmp_path / "baixado.xlsx")
    assert sp.download(2024, cache_dir=tmp_path) == tmp_path / "baixado.xlsx"
    assert len(calls) == 2


def test_download_uses_the_cached_copy_when_the_page_is_unreachable(monkeypatch, tmp_path, caplog):
    cached = tmp_path / "abc1234567_itbi_sp_2024.xlsx"
    cached.write_bytes(b"xlsx")

    def unreachable():
        raise requests.ConnectionError("sem rede")

    monkeypatch.setattr(sp, "list_files", unreachable)
    with caplog.at_level(logging.WARNING, logger="dadosimob"):
        assert sp.download(2024, cache_dir=tmp_path) == cached
    assert "cache" in caplog.text


def test_download_without_page_or_cache_raises(monkeypatch, tmp_path):
    monkeypatch.setattr(sp, "list_files", lambda: [])
    with pytest.raises(ValueError, match="2024"):
        sp.download(2024, cache_dir=tmp_path)


@pytest.mark.network
def test_list_files_live():
    files = sp.list_files()
    assert any(f.year == 2024 for f in files)


def test_cli(workbook, tmp_path, capsys):
    from dadosimob.cli import main

    out = tmp_path / "saida.csv"
    assert main(["itbi-sp", "--arquivo", str(workbook), "--limpar", "--apenas-vendas", "-o", str(out)]) == 0
    df = pd.read_csv(out)
    assert len(df) == 4  # drops the too cheap shop and the donation
    assert "linhas salvas" in capsys.readouterr().out
